from docx import Document
from docx.shared import Pt, RGBColor
from pathlib import Path

root=Path(__file__).parent
d=Document(root/'current-before-technologies.docx')
t=d.tables[1]
assert [[c.text for c in r.cells] for r in t.rows][0]==['Component','Technology']
rows=[
 ('Programming and markup', 'JavaScript (ES modules), JSX, HTML5 and plain CSS for application logic, page structure and styling.'),
 ('Frontend UI library', 'React 19.3.1 and React DOM 19.3.1 for reusable components and interactive resident and SDP screens.'),
 ('Development and build tools', 'Vite 8.3.4 with the React plugin for the development server and production builds.'),
 ('Testing and code quality', 'Vitest for unit tests, Playwright for browser tests and ESLint for code checks.'),
 ('Planned 2D world framework', 'Phaser 3.90.0 for the pixel-art map and world interactions; planned and not yet installed in the current frontend shell.')
]
for row,content in zip(t.rows[1:],rows):
    for cell,text in zip(row.cells,content):
        p=cell.paragraphs[0]
        p.text=text
        for extra in list(cell.paragraphs[1:]):extra._p.getparent().remove(extra._p)
        p.paragraph_format.line_spacing=1.15
        for r in p.runs:
            r.font.name='Times New Roman';r.font.size=Pt(12)
            r.font.color.rgb=RGBColor.from_string('6C1A21')
        p.paragraph_format.keep_together=True
assert all(c.text.strip() for row in t.rows[1:] for c in row.cells)
d.save(root/'KoBi-report-technologies.docx')
print('Updated five technology rows; preserved other report content.')
