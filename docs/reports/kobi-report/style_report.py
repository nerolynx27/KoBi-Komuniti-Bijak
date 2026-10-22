from pathlib import Path
import re, io
from collections import defaultdict
from copy import deepcopy
import pdfplumber, pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).parent
WINE='6C1A21'; GOLD='CEA257'; SAND='D1BF9A'; VANILLA='E5E0CD'

def evidence(n, page_no, pattern, left=False):
    src=ROOT/f'source-{n}.pdf'
    with pdfplumber.open(src) as pdf:
        page=pdf.pages[page_no]
        scope=page.crop((0,0,page.width/2,page.height)) if left else page.crop((165,0,page.width,page.height))
        hits=scope.search('(?s)'+pattern,regex=True,case=False,x_tolerance=1)
        if not hits: raise RuntimeError(f'No evidence match {n}')
        chars=hits[0]['chars']
        lines=defaultdict(list)
        for ch in chars: lines[round(ch['top'],1)].append(ch)
        buf=io.BytesIO(); c=canvas.Canvas(buf,pagesize=(page.width,page.height))
        c.setFillColorRGB(1,0.85,0); c.setFillAlpha(.35)
        for row in lines.values():
            x0=min(ch['x0'] for ch in row); x1=max(ch['x1'] for ch in row)
            top=min(ch['top'] for ch in row); bottom=max(ch['bottom'] for ch in row)
            c.rect(x0-1,page.height-bottom-1,x1-x0+2,bottom-top+2,fill=1,stroke=0)
        c.save(); buf.seek(0)
        p=PdfReader(src).pages[page_no]; p.merge_page(PdfReader(buf).pages[0])
        out=PdfWriter();out.add_page(p)
        target=ROOT/f'evidence-{n}-highlighted.pdf'
        with target.open('wb') as f: out.write(f)
        render=pdfium.PdfDocument(str(target)); render[0].render(scale=2).to_pil().save(ROOT/f'evidence-{n}.png')
        print('Evidence',n,'page',page_no+1,'lines',len(lines))

evidence(1,1,r'Most\s+of the included studies considered citizens.*?support city management\.')
evidence(2,5,r'Information provided by the apps helps to remove.*?how to recycle appropriately\.',True)

d=Document(ROOT/'original.docx')
orig=list(d.paragraphs)
# Remove manual blank spacers and page breaks; preserve the existing cover image.
for p in list(d.paragraphs):
    for br in list(p._p.xpath('.//w:br[@w:type="page"]')): br.getparent().remove(br)
    if not p.text.strip() and not p._p.xpath('.//w:drawing'):
        p._p.getparent().remove(p._p)

for st in d.styles:
    if st.type==1 or st.type==2:
        st.font.name='Times New Roman'; st.font.size=Pt(12); st.font.color.rgb=RGBColor.from_string(WINE)
        if st.type==1:
            st.paragraph_format.line_spacing=1.15
            st.paragraph_format.space_after=Pt(6)
            st.paragraph_format.space_before=Pt(0)
            st.paragraph_format.left_indent=Inches(0)
            st.paragraph_format.first_line_indent=Inches(0)
            st.paragraph_format.keep_with_next=False
            st.paragraph_format.page_break_before=False

sec=d.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.different_first_page_header_footer=True
for tag in ('w:pgNumType','w:docGrid'):
    for e in list(sec._sectPr.findall(qn(tag))):sec._sectPr.remove(e)

def format_run(r,font='Times New Roman',size=12,bold=None,color=WINE):
    r.font.name=font;r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
    rp=r._r.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),font)
    for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+a),None)
    if bold is not None:r.bold=bold
    r.font.all_caps=False
    r.font.highlight_color=None

def normal(p):
    p.style=d.styles['Normal'];f=p.paragraph_format
    f.line_spacing=1.15;f.space_after=Pt(6);f.space_before=Pt(0)
    f.left_indent=Inches(0);f.right_indent=Inches(0);f.first_line_indent=Inches(0)
    f.keep_with_next=False;f.keep_together=False;f.page_break_before=False
    pp=p._p.get_or_add_pPr()
    for tag in ['w:numPr','w:pBdr','w:shd','w:tabs','w:ind','w:rPr']:
        for e in list(pp.findall(qn(tag))):pp.remove(e)
    for r in p.runs:format_run(r)

def heading(p,level=1):
    normal(p);p.style=d.styles[f'Heading {level}']
    f=p.paragraph_format;f.space_before=Pt(16 if level==1 else 9);f.space_after=Pt(7);f.keep_with_next=True
    for r in p.runs:format_run(r,'Bungee',16 if level==1 else 12,False)

heading_indices={21,37,55,63,85,98}
sub_indices={39,44,49,57,65,73,87,100,103}
for i,p in enumerate(orig):
    if p._p.getparent() is None:continue
    normal(p)
    if i in heading_indices:heading(p)
    elif i in sub_indices:heading(p,2)
    elif i in {41,46,51}:
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.bold=True

# Keep the cover's content and image, with deliberate typography and spacing.
for p in orig[:21]:
    if p._p.getparent() is not None:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in orig[1].runs:format_run(r,'Bungee',26,False)
orig[1].style=d.styles['Title'];orig[1].paragraph_format.space_before=Pt(25);orig[1].paragraph_format.space_after=Pt(12)
for r in orig[2].runs:format_run(r,'Yanone Kaffeesatz',20,False)
orig[2].paragraph_format.space_after=Pt(28)
orig[0].paragraph_format.space_before=Pt(35)
orig[21].paragraph_format.page_break_before=True
orig[55].paragraph_format.page_break_before=True
orig[85].paragraph_format.page_break_before=True
orig[23].style=d.styles['Normal']

def after(p,text):
    np=d.add_paragraph(text);p._p.addnext(np._p);normal(np);return np

c1=after(orig[42], 'Research on citizen participation applications identifies reporting problems with urban infrastructure maintenance as a common use of these tools (Bastos et al., 2022).')
c2=after(orig[52], 'Dursun et al. (2024) explain that recycling apps can address knowledge barriers by helping users recognise recyclable materials and understand appropriate recycling methods. This supports KoBi’s provision of accepted-material and collection information.')

# Native numbered lists retain the order of the two workflows.
def add_numbering(paragraphs):
    root=d.part.numbering_part.element
    aid=max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))]+[0])+1
    nid=max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(aid))
    ns=OxmlElement('w:nsid');ns.set(qn('w:val'),f'{aid+100000:08X}');abstract.append(ns)
    mt=OxmlElement('w:multiLevelType');mt.set(qn('w:val'),'singleLevel');abstract.append(mt)
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);lvl.append(el)
    abstract.append(lvl)
    firstnum=root.find(qn('w:num'))
    if firstnum is not None:firstnum.addprevious(abstract)
    else:root.append(abstract)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));a=OxmlElement('w:abstractNumId');a.set(qn('w:val'),str(aid));num.append(a)
    override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0');start=OxmlElement('w:startOverride');start.set(qn('w:val'),'1');override.append(start);num.append(override);root.append(num)
    for p in paragraphs:
        pp=p._p.get_or_add_pPr();np=OxmlElement('w:numPr')
        for tag,val in [('ilvl','0'),('numId',str(nid))]:
            el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);np.append(el)
        pp.append(np);p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22)
add_numbering(orig[67:72]);add_numbering(orig[75:79])

# Preserve native tables and empty technology cells; only change presentation.
for table in d.tables:
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr()
        for el in list(trpr.findall(qn('w:trHeight'))):trpr.remove(el)
        cs=OxmlElement('w:cantSplit');trpr.append(cs)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(2 if ci==0 else 4.67)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text=cell.text.strip()
            cell.text=text
            for p in cell.paragraphs:
                normal(p);p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                if ri==0:
                    for r in p.runs:format_run(r,'Bungee',12,False,WINE)
            tcpr=cell._tc.get_or_add_tcPr()
            for tag in ['w:shd','w:tcBorders']:
                for el in list(tcpr.findall(qn(tag))):tcpr.remove(el)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),GOLD if ri==0 else (VANILLA if ri%2 else SAND));tcpr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),WINE);borders.append(el)
            tcpr.append(borders)
    # Reset any old theme table borders and banding.
    table.style=d.styles['Table Grid']

# References use APA 7 sentence case and a half-inch hanging indent.
p=d.add_paragraph('8. References');heading(p);p.paragraph_format.page_break_before=True
refs=[
('Bastos, D., Fernández-Caballero, A., Pereira, A., & Rocha, N. P. (2022). Smart city applications to promote citizen participation in city management and governance: A systematic review. ', 'Informatics, 9', '(4), Article 89. https://doi.org/10.3390/informatics9040089'),
('Dursun, İ., Tümer, E., & Yürüyen Kılıç, H. (2024). Roles of mobile applications in removing barriers to individual recycling: Case of Türkiye. ', 'Environmental Research and Technology, 7', '(1), 97–107. https://doi.org/10.35208/ert.1317616')]
for start,journal,end in refs:
    p=d.add_paragraph();normal(p)
    p.add_run(start);p.add_run(journal).italic=True;p.add_run(end)
    for r in p.runs:format_run(r)
    p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
    p.paragraph_format.space_after=Pt(12)

for n,title,caption in [
    (1,'Appendix A — Citation evidence','Bastos et al. (2022), article page 1 (PDF page 2). The highlighted abstract passage supports the infrastructure-reporting statement in Section 3.'),
    (2,'Appendix B — Citation evidence','Dursun et al. (2024), article page 102 (PDF page 6). The highlighted passage supports the recycling-information statement in Section 3.')]:
    p=d.add_paragraph(title);heading(p);p.paragraph_format.page_break_before=True
    p=d.add_paragraph(caption);normal(p)
    p=d.add_paragraph();normal(p);p.add_run().add_picture(str(ROOT/f'evidence-{n}.png'),width=Inches(6.35));p.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Refresh footer typography while preserving the existing page field.
for section in d.sections:
    for p in section.footer.paragraphs:
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:format_run(r,'Times New Roman',10,False)
d.save(ROOT/'KoBi-report-styled.docx')
print('Saved styled report')
